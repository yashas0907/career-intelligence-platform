"""Resume fixer endpoints: fix in-place + downloadable DOCX/TXT."""

from __future__ import annotations

import io
import logging

from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from app.core.database import get_db
from app.models.entities import Job, Resume
from app.services.resume_fixer import fix_resume

router = APIRouter(prefix="/fix", tags=["fixer"])
logger = logging.getLogger("app.api.fix")


class FixResponse(BaseModel):
    resume_id: str
    fixed_text: str
    changes: list[dict[str, str]]
    llm_used: bool


@router.post("/{resume_id}", response_model=FixResponse)
def fix_resume_endpoint(
    resume_id: str,
    job_id: str | None = None,
    db: Session = Depends(get_db),
) -> FixResponse:
    resume = db.get(Resume, resume_id)
    if not resume:
        raise HTTPException(status_code=404, detail="Resume not found.")
    parsed_jd = None
    if job_id:
        job = db.get(Job, job_id)
        if not job:
            raise HTTPException(status_code=404, detail="Job not found.")
        parsed_jd = dict(job.parsed or {})

        # annotate the top missing required skills for the fixer's hints
        cand = set((resume.profile.get("skills") or {}).keys())
        missing = [s for s in (parsed_jd.get("required_skills") or []) if s not in cand]
        parsed_jd["missing_top"] = missing[:8]

    result = fix_resume(resume.profile, resume.raw_text, parsed_jd)
    logger.info("Fixed resume=%s changes=%d llm=%s", resume_id, len(result.changes), result.llm_used)
    return FixResponse(
        resume_id=resume_id,
        fixed_text=result.fixed_text,
        changes=result.changes,
        llm_used=result.llm_used,
    )


@router.get("/{resume_id}/download")
def download_fixed(
    resume_id: str,
    format: str = "txt",
    job_id: str | None = None,
    db: Session = Depends(get_db),
) -> StreamingResponse:
    """Download the fixed resume as TXT or DOCX (real Word file)."""
    if format not in {"txt", "docx"}:
        raise HTTPException(status_code=422, detail="format must be txt or docx")

    resume = db.get(Resume, resume_id)
    if not resume:
        raise HTTPException(status_code=404, detail="Resume not found.")
    parsed_jd = None
    if job_id:
        job = db.get(Job, job_id)
        if job:
            parsed_jd = dict(job.parsed or {})
            cand = set((resume.profile.get("skills") or {}).keys())
            parsed_jd["missing_top"] = [s for s in (parsed_jd.get("required_skills") or []) if s not in cand][:8]

    result = fix_resume(resume.profile, resume.raw_text, parsed_jd)
    base_name = (resume.filename or "resume").rsplit(".", 1)[0]

    if format == "txt":
        content = result.fixed_text
        return StreamingResponse(
            iter([content]),
            media_type="text/plain; charset=utf-8",
            headers={"Content-Disposition": f'attachment; filename="{base_name}-fixed.txt"'},
        )

    # DOCX
    from docx import Document

    doc = Document()
    for block in result.fixed_text.split("\n\n"):
        lines = block.splitlines()
        if not lines:
            continue
        head = lines[0].strip()
        if head.isupper() and len(head) < 40:  # section heading
            doc.add_heading(head.title(), level=1)
            for line in lines[1:]:
                if line.strip():
                    doc.add_paragraph(line.strip())
        else:
            for i, line in enumerate(lines):
                text = line.strip()
                if not text:
                    continue
                if text.startswith("- "):
                    doc.add_paragraph(text[2:], style="List Bullet")
                elif i == 0:
                    doc.add_heading(text, level=2) if len(text) < 60 else doc.add_paragraph(text)
                else:
                    doc.add_paragraph(text)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        headers={"Content-Disposition": f'attachment; filename="{base_name}-fixed.docx"'},
    )

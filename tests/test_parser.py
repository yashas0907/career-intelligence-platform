"""Tests: document parser validation + robustness."""

import pytest

from app.services.document_parser import DocumentParseError, parse_document


class TestValidation:
    def test_rejects_unsupported_extension(self):
        with pytest.raises(DocumentParseError, match="Unsupported file type"):
            parse_document("virus.exe", b"MZ\x90\x00")

    def test_rejects_empty_file(self):
        with pytest.raises(DocumentParseError, match="empty"):
            parse_document("empty.pdf", b"")

    def test_rejects_oversized_file(self, monkeypatch):
        monkeypatch.setattr("app.services.document_parser.MAX_BYTES", 100)
        with pytest.raises(DocumentParseError, match="too large"):
            parse_document("big.txt", b"x" * 1000)

    def test_rejects_contentless_pdf(self, monkeypatch):
        # simulate pdfminer returning nothing
        monkeypatch.setattr("app.services.document_parser._parse_pdf", lambda data: "")
        with pytest.raises(DocumentParseError, match="No readable text"):
            parse_document("scan.pdf", b"%PDF-1.4 fake")

    def test_no_extension_rejected(self):
        with pytest.raises(DocumentParseError):
            parse_document("README", b"hello world")


class TestTxtParsing:
    def test_parses_utf8(self):
        doc = parse_document("resume.txt", "Alice\nPython developer".encode("utf-8"))
        assert doc.text == "Alice\nPython developer"
        assert doc.method == "txt"
        assert doc.pages == 1

    def test_parses_latin1_fallback(self):
        doc = parse_document("resume.txt", "José Pythón developer".encode("latin-1"))
        assert "developer" in doc.text

    def test_control_chars_cleaned(self):
        doc = parse_document("r.txt", b"hello\x00\x07world\x1f!")
        assert "\x00" not in doc.text

    def test_real_txt_resume(self, sample_texts):
        doc = parse_document("resume.txt", sample_texts["resume"].encode())
        assert "PyTorch" in doc.text
        assert doc.char_count > 500


class TestDocxParsing:
    def test_parses_docx_with_paragraphs_and_tables(self):
        from docx import Document as DocxDocument

        import io

        d = DocxDocument()
        d.add_paragraph("Alice Smith")
        d.add_paragraph("Skills: Python, PyTorch")
        t = d.add_table(rows=1, cols=2)
        t.cell(0, 0).text = "AWS"
        t.cell(0, 1).text = "Docker"
        buf = io.BytesIO()
        d.save(buf)

        doc = parse_document("resume.docx", buf.getvalue())
        assert "Alice Smith" in doc.text
        assert "AWS" in doc.text  # from table
        assert doc.method == "docx"

    def test_rejects_fake_docx(self):
        with pytest.raises(DocumentParseError, match="DOCX"):
            parse_document("fake.docx", b"this is not a real docx zip")

    def test_empty_docx(self):
        from docx import Document as DocxDocument

        import io

        d = DocxDocument()
        buf = io.BytesIO()
        d.save(buf)
        with pytest.raises(DocumentParseError, match="No readable text"):
            parse_document("empty.docx", buf.getvalue())


class TestPdfParsing:
    def test_real_pdf(self):
        # Build a minimal real PDF with one text page.
        import textwrap

        content = textwrap.dedent(
            """\
            BT /F1 12 Tf 50 700 Td (Alice Smith - Python Developer) Tj ET
            BT /F1 10 Tf 50 680 Td (Skills: Python PyTorch Docker) Tj ET
            endstream"""
        )
        stream = content.encode("latin-1")
        objs = []
        objs.append(b"<< /Type /Catalog /Pages 2 0 R >>")
        objs.append(b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
        objs.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>"
        )
        objs.append(b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream")
        objs.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

        out = b"%PDF-1.4\n"
        offsets = []
        for i, obj in enumerate(objs, start=1):
            offsets.append(len(out))
            out += f"{i} 0 obj\n".encode() + obj + b"\nendobj\n"
        xref_pos = len(out)
        out += b"xref\n0 " + str(len(objs) + 1).encode() + b"\n"
        out += b"0000000000 65535 f \n"
        for off in offsets:
            out += f"{off:010d} 00000 n \n".encode()
        out += (
            b"trailer\n<< /Size " + str(len(objs) + 1).encode() + b" /Root 1 0 R >>\nstartxref\n"
            + str(xref_pos).encode() + b"\n%%EOF"
        )

        doc = parse_document("resume.pdf", out)
        assert "Alice Smith" in doc.text
        assert doc.method == "pdfminer"

    def test_corrupted_pdf(self):
        with pytest.raises(DocumentParseError):
            parse_document("bad.pdf", b"%PDF-1.7 " + b"\x00\x01\x02garbage" * 40)

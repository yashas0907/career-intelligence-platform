"""Document ingestion: PDF / DOCX / TXT -> clean text with structural hints.

Robustness rules:
* Unsupported extensions and empty/oversized files are rejected with clear errors.
* Malformed binaries raise DocumentParseError (never a 500 with a stack trace).
* PDFs are parsed with pdfminer.six (pure python); DOCX via python-docx.
* Whitespace/control chars normalized; page markers preserved as \f for chunking.
"""

from __future__ import annotations

import logging
import re
from dataclasses import dataclass

from app.core.config import settings
from app.services.skills import normalizer

logger = logging.getLogger("app.parser")

MAX_BYTES = settings.max_upload_mb * 1024 * 1024


class DocumentParseError(Exception):
    """Raised for any user-facing ingestion failure."""


@dataclass
class ParsedDocument:
    text: str
    pages: int
    method: str  # pdfminer | docx | txt
    char_count: int
    warnings: list[str]


_CTRL_RE = re.compile(r"[\x00-\x08\x0b\x0e-\x1f\x7f]")


def _clean(text: str) -> str:
    text = _CTRL_RE.sub(" ", text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def validate_extension(filename: str) -> str:
    ext = (filename.rsplit(".", 1)[-1] if "." in filename else "").lower()
    if ext not in settings.allowed_extensions:
        raise DocumentParseError(
            f"Unsupported file type '.{ext}'. Allowed: {', '.join(settings.allowed_extensions)}"
        )
    return ext


def validate_size(size_bytes: int) -> None:
    if size_bytes <= 0:
        raise DocumentParseError("The uploaded file is empty.")
    if size_bytes > MAX_BYTES:
        raise DocumentParseError(
            f"File too large ({size_bytes / 1024 / 1024:.1f} MB). Maximum is {settings.max_upload_mb} MB."
        )


def _parse_pdf(data: bytes) -> str:
    import io

    from pdfminer.high_level import extract_text

    try:
        return extract_text(io.BytesIO(data)) or ""
    except Exception as exc:  # noqa: BLE001  (pdfminer raises assorted exception types)
        if isinstance(exc, (ValueError, OSError)) and "pdf" in str(exc).lower():
            raise DocumentParseError("The PDF appears to be corrupted or is not a valid PDF.") from exc
        logger.exception("pdfminer failed: %s", exc)
        raise DocumentParseError("Failed to read the PDF file.") from exc


def _parse_docx(data: bytes) -> str:
    try:
        import io

        from docx import Document as DocxDocument

        doc = DocxDocument(io.BytesIO(data))
        parts: list[str] = []
        for para in doc.paragraphs:
            para_text = para.text.strip()
            if para_text:
                parts.append(para_text)
        # tables are common in resumes (skills matrices) — keep them
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts)
    except Exception as exc:  # noqa: BLE001
        logger.exception("python-docx failed: %s", exc)
        raise DocumentParseError(
            "Failed to read the DOCX file. It may be corrupted or not a real .docx (e.g. renamed .doc)."
        ) from exc


def _parse_txt(data: bytes) -> str:
    for enc in ("utf-8", "utf-8-sig", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise DocumentParseError("Could not decode the text file (unsupported encoding).")


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    ext = validate_extension(filename)
    validate_size(len(data))

    warnings: list[str] = []

    if ext == "pdf":
        text = _parse_pdf(data)
        method = "pdfminer"
    elif ext == "docx":
        text = _parse_docx(data)
        method = "docx"
    else:
        text = _parse_txt(data)
        method = "txt"

    text = _clean(text)
    if not text:
        raise DocumentParseError(
            "No readable text found in the document. Scanned/image-only PDFs are not supported."
        )

    if len(text) > settings.max_resume_chars:
        warnings.append(
            f"Document truncated to {settings.max_resume_chars} characters for analysis."
        )
        text = text[: settings.max_resume_chars]

    # sanity: a resume should mention at least some taxonomy terms
    hits = normalizer.extract(text[:3000])
    if len(hits) < 2:
        warnings.append(
            "Very few recognizable skills found — the document may not be a resume or uses unusual terminology."
        )

    pages = max(1, text.count("\f") + 1)
    logger.info(
        "Parsed %s (%s): %d chars, %d pages, warnings=%d",
        filename, method, len(text), pages, len(warnings),
    )
    return ParsedDocument(text=text, pages=pages, method=method, char_count=len(text), warnings=warnings)
